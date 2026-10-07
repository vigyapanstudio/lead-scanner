"""Builds the colour-coded Excel report from scan results."""
from __future__ import annotations

import datetime as dt
import io

import pandas as pd
from brand import CLIENT_NAME, POWERED_BY
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "0B2E2C"
ACCENT = "FFB72B"
RED, AMBER, GREEN = "F8D7DA", "FFF1CC", "D8F0DF"
RED_T, AMBER_T, GREEN_T = "8A1C24", "7A5300", "1E6B37"
THIN = Side(style="thin", color="D9DCE3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MAIN_COLS = [
    "Lead Priority", "Event", "Company", "Contact Person", "Phone", "Email", "City", "Category",
    "Opportunities", "Pitch Note", "Pitch (Hinglish)", "WhatsApp Link", "WhatsApp Message",
    "Coverage Prospect", "Coverage Pitch", "Data Check",
    "Website", "Website Checked", "Website Status", "Website Score", "Website Age",
    "SEO Score", "SEO Status", "Instagram Handle", "Instagram Status",
    "Instagram Followers", "Instagram Posts", "Instagram Last Post", "Instagram Engagement %",
    "Other Socials Found", "Facebook", "LinkedIn", "YouTube", "Google Maps", "Google Rating", "Google Reviews",
    "Google Maps Link", "Email (from website)", "Phone (from website)",
    "Website Issues", "SEO Issues", "Instagram Issues", "Social Media Issues",
    "HTTPS Secure", "Mobile Friendly", "Load Time (s)", "Copyright Year", "Domain Registered",
    "SSL Days Left", "Built With", "Old Technology", "Mobile Speed (Google)",
    "Page Title", "Meta Description", "Notes",
]
WIDTHS = {"Lead Priority": 11, "Company": 26, "Contact Person": 20, "Phone": 15, "Email": 28,
          "Opportunities": 30, "Pitch Note": 55, "Website": 26, "Website Checked": 30,
          "Website Status": 24, "Website Age": 26, "SEO Status": 15, "Instagram Status": 24,
          "Website Issues": 55, "SEO Issues": 50, "Instagram Issues": 45, "Page Title": 35,
          "Meta Description": 40, "Notes": 40, "Old Technology": 25, "Pitch (Hinglish)": 55,
          "WhatsApp Link": 18, "WhatsApp Message": 60, "Coverage Pitch": 45, "Data Check": 24,
          "Social Media Issues": 40, "Facebook": 22, "LinkedIn": 22, "YouTube": 22, "Google Maps Link": 22}
CONTACT_COLS = ["Lead Priority", "Company", "Contact Person", "Phone", "Email", "City",
                "Website", "WhatsApp Link", "Pitch Note"]

LISTS = [
    ("No Website", lambda d: d["Opportunities"].str.contains("New Website"),
     ["Website Status", "Website Issues"]),
    ("Website Redesign", lambda d: d["Opportunities"].str.contains("Website Redesign|Website Upgrade"),
     ["Website Checked", "Website Status", "Website Score", "Website Age", "Website Issues"]),
    ("Needs SEO", lambda d: d["Opportunities"].str.contains("SEO") & ~d["Opportunities"].str.contains("New Website"),
     ["Website Checked", "SEO Score", "SEO Issues"]),
    ("Instagram Leads", lambda d: d["Opportunities"].str.contains("Instagram"),
     ["Instagram Handle", "Instagram Status", "Instagram Issues"]),
    ("Google Maps Leads", lambda d: d["Opportunities"].str.contains("Google Business"),
     ["Google Maps", "Google Rating", "Google Reviews", "Google Maps Link"]),
    ("Coverage Prospects", lambda d: d.get("Coverage Prospect", pd.Series([""] * len(d), index=d.index)) == "Yes",
     ["Website Checked", "Instagram Handle", "Coverage Pitch"]),
]


def _avg(series):
    m = pd.to_numeric(series, errors="coerce").mean()
    return "-" if pd.isna(m) else round(m)


def _header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def _write_table(ws, df: pd.DataFrame, start_row=1):
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        ws.cell(row=start_row, column=j, value=c)
    _header(ws, start_row, len(cols))
    for i, rec in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(rec, 1):
            if v is None or (isinstance(v, float) and pd.isna(v)):
                v = None
            col_name = cols[j - 1]
            if col_name in ("WhatsApp Link", "Google Maps Link", "Facebook", "LinkedIn", "YouTube") \
                    and isinstance(v, str) and v.startswith("http"):
                cell = ws.cell(row=i, column=j, value="Open WhatsApp" if col_name == "WhatsApp Link" else v)
                cell.hyperlink = v
                cell.font = Font(color="0B6E4F", underline="single", bold=col_name == "WhatsApp Link")
            else:
                cell = ws.cell(row=i, column=j, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(c, 14 if len(c) < 14 else len(c) + 2)
    last = start_row + len(df)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=3)
    if len(df):
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{last}"
    # colour rules
    for j, c in enumerate(cols, 1):
        L = get_column_letter(j)
        rng = f"{L}{start_row + 1}:{L}{max(last, start_row + 1)}"
        if c in ("Website Score", "SEO Score", "Mobile Speed (Google)"):
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["50"],
                                          fill=PatternFill("solid", fgColor=RED), font=Font(color=RED_T, bold=True)))
            ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["50", "79"],
                                          fill=PatternFill("solid", fgColor=AMBER), font=Font(color=AMBER_T, bold=True)))
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["80"],
                                          fill=PatternFill("solid", fgColor=GREEN), font=Font(color=GREEN_T, bold=True)))
        if c == "Lead Priority":
            for val, fill, font in (("HOT", RED, RED_T), ("WARM", AMBER, AMBER_T), ("LOW", GREEN, GREEN_T)):
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{val}"'],
                                              fill=PatternFill("solid", fgColor=fill), font=Font(color=font, bold=True)))
    return last


def build_report(results: pd.DataFrame, source_name: str = "", removed: pd.DataFrame | None = None) -> bytes:
    df = results.copy()
    order = {"HOT": 0, "WARM": 1, "CHECK": 2, "LOW": 3}
    df["_o"] = df["Lead Priority"].map(order).fillna(3)
    df = df.sort_values(["_o", "Website Score"], na_position="first").drop(columns="_o")
    cols = [c for c in MAIN_COLS if c in df.columns] + [c for c in df.columns if c.startswith("[Original]")]
    # drop contact fields that were not in the source at all
    cols = [c for c in cols if not (c in ("City", "Category", "Contact Person") and (df[c].astype(str).str.strip() == "").all())]
    df = df[cols]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws["B2"] = f"{CLIENT_NAME.upper()}  ·  Lead Scanner Report"
    ws["B2"].font = Font(size=18, bold=True, color=NAVY)
    ws["B3"] = f"Source: {source_name or 'database'}   ·   Scanned on {dt.datetime.now():%d %b %Y, %I:%M %p}   ·   Powered by {POWERED_BY}"
    ws["B3"].font = Font(size=10, color="6B7280")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14

    total = len(df)
    opp = df["Opportunities"].fillna("")

    def cnt(pattern):
        return int(opp.str.contains(pattern).sum())

    rows = [
        ("Total contacts scanned", total),
        ("HOT leads (need 2+ big services)", int((df["Lead Priority"] == "HOT").sum())),
        ("WARM leads (need 1-2 services)", int((df["Lead Priority"] == "WARM").sum())),
        ("Digitally healthy (LOW)", int((df["Lead Priority"] == "LOW").sum())),
        ("Check manually (site blocked our scan)", int((df["Lead Priority"] == "CHECK").sum())),
        (None, None),
        ("Need a NEW website", cnt("New Website")),
        ("Need website REDESIGN", cnt("Website Redesign")),
        ("Need website UPGRADE", cnt("Website Upgrade")),
        ("Need SEO work", cnt("SEO")),
        ("Need Instagram page setup", cnt("Instagram Page Setup")),
        ("Need Instagram management", cnt("Instagram Management")),
        (None, None),
        ("Average website score", _avg(df["Website Score"])),
        ("Average SEO score (sites that open)", _avg(df["SEO Score"])),
    ]
    r0 = 5
    ws.cell(row=r0, column=2, value="What we found")
    ws.cell(row=r0, column=3, value="Count")
    ws.cell(row=r0, column=4, value="% of list")
    _header(ws, r0, 4)
    ws.cell(row=r0, column=1).fill = PatternFill(fill_type=None)
    ws.cell(row=r0, column=1).border = Border()
    r = r0
    for label, val in rows:
        r += 1
        if label is None:
            continue
        ws.cell(row=r, column=2, value=label).border = BORDER
        c = ws.cell(row=r, column=3, value=val)
        c.border = BORDER
        c.font = Font(bold=True)
        if "score" not in label and label != "Total contacts scanned" and total:
            p = ws.cell(row=r, column=4, value=val / total)
            p.number_format = "0%"
            p.border = BORDER
        if label.startswith("HOT"):
            ws.cell(row=r, column=2).font = Font(bold=True, color=RED_T)

    # chart of services needed
    chart_rows = [(l, v) for l, v in rows[6:12]]
    ws.cell(row=r0, column=7, value="Service")
    ws.cell(row=r0, column=8, value="Leads")
    for i, (l, v) in enumerate(chart_rows, 1):
        ws.cell(row=r0 + i, column=7, value=l.replace("Need ", ""))
        ws.cell(row=r0 + i, column=8, value=v)
    ws.column_dimensions["G"].width = 28
    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = "Services these leads need"
    ch.legend = None
    ch.add_data(Reference(ws, min_col=8, min_row=r0, max_row=r0 + len(chart_rows)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=7, min_row=r0 + 1, max_row=r0 + len(chart_rows)))
    ch.series[0].graphicalProperties.solidFill = ACCENT
    ch.height, ch.width = 8, 16
    ws.add_chart(ch, "J5")

    ws.cell(row=r + 2, column=2, value="Tabs in this file").font = Font(bold=True, color=NAVY, size=12)
    tabs = [("All Leads", "Every contact with scores, issues and a ready pitch line (HOT first)"),
            ("No Website", "Companies with no website / broken / parked site"),
            ("Website Redesign", "Websites that are old, not mobile-friendly, insecure or slow"),
            ("Needs SEO", "Working websites that Google can't read well"),
            ("Instagram Leads", "No Instagram page, or an Instagram that needs work"),
            ("Google Maps Leads", "Not on Google Maps, or very few reviews"),
            ("Coverage Prospects", "Already strong online - pitch event coverage / podcast"),
            ("Cleanup Log", "Duplicate contacts removed before scanning"),
            ("How to Read", "What each score and column means")]
    for i, (t, d) in enumerate(tabs, r + 3):
        ws.cell(row=i, column=2, value=t).font = Font(bold=True)
        ws.cell(row=i, column=3, value=d)

    # All leads
    wa = wb.create_sheet("All Leads")
    _write_table(wa, df)

    for name, cond, extra in LISTS:
        sub = df[cond(df.fillna(""))] if len(df) else df
        keep = [c for c in CONTACT_COLS if c in df.columns] + [c for c in extra if c in df.columns]
        wsx = wb.create_sheet(name)
        _write_table(wsx, sub[keep])

    if removed is not None and len(removed):
        wl = wb.create_sheet("Cleanup Log")
        wl["A1"] = "These duplicate contacts were removed before scanning (same phone, email or company):"
        wl["A1"].font = Font(bold=True, color=NAVY)
        _write_table(wl, removed.drop(columns=[c for c in removed.columns if c == "Data Check"]), start_row=3)

    wh = wb.create_sheet("How to Read")
    wh.column_dimensions["A"].width = 26
    wh.column_dimensions["B"].width = 100
    guide = [
        ("Lead Priority", "HOT = needs 2+ big services (e.g. no website AND no Instagram). WARM = needs 1-2 services. LOW = already in good digital shape."),
        ("Website Status", "No Website · Broken / Not Opening · Parked / Empty / Error · Outdated - Needs Redesign (score under 60) · Needs Upgrade (60-79) · Good / Modern (80+)."),
        ("Website Score (0-100)", "Starts at 100. Points are cut for: no HTTPS security, not mobile-friendly, old footer year, old technology (Flash, old jQuery/WordPress, table layouts), slow loading, no contact buttons, DIY builder."),
        ("Website Age", "Based on the latest year seen in the footer (©) or server 'last modified' date. 'Old' = no sign of update in 4+ years."),
        ("Domain Registered", "Year the domain name was first bought (public record). Tells how long the business has been online."),
        ("SEO Score (0-100)", "Checks the basics Google needs: page title, description, main heading, image alt text, sitemap, robots.txt, social share tags, business schema, analytics, enough text. Under 70 = Needs SEO Work."),
        ("Instagram Status", "Found from the list or from links on the website. Detailed checks (followers, last post, bio, link in bio, business account, engagement) only run if an Apify token was added."),
        ("Opportunities", "The services you can pitch: New Website, Website Redesign/Upgrade, SEO, Instagram Page Setup, Instagram Management, Social Media Setup."),
        ("Pitch Note", "A one-line, plain-English reason to call this lead (Hinglish version next to it). Edit before using."),
        ("WhatsApp Link", "Click to open WhatsApp with a ready, personalised message to this contact. Check it, then press send."),
        ("Coverage Prospect", "Already strong online, so likely to have a marketing budget - pitch event coverage or a podcast feature."),
        ("Google Maps", "Only filled in when the Google Maps check is switched on. Shows rating and number of reviews."),
        ("Data Check", "Problems found in your list before scanning, e.g. a phone number that is too short."),
        ("Colours", "Red = poor (under 50), Yellow = average (50-79), Green = good (80+)."),
        ("Limits", "This is an automated first check of the HOME PAGE only. Always open the site yourself before pitching. Some sites block automated visitors and may show as 'not opening'."),
    ]
    wh["A1"] = "How to read this report"
    wh["A1"].font = Font(size=14, bold=True, color=NAVY)
    for i, (k, v) in enumerate(guide, 3):
        wh.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = wh.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        wh.row_dimensions[i].height = 32

    for s in wb.worksheets:
        s.sheet_properties.tabColor = ACCENT if s.title in ("Summary", "All Leads") else NAVY
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
